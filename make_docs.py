import os, json
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from pptx import Presentation
from pptx.util import Inches as PInches, Pt as PPt
from pptx.enum.text import PP_ALIGN

ROOT=os.path.dirname(os.path.abspath(__file__)); OUT=os.path.join(ROOT,'outputs'); PLOT=os.path.join(OUT,'plots'); DOCS=os.path.join(ROOT,'docs')
os.makedirs(DOCS,exist_ok=True)
with open(os.path.join(OUT,'project_summary.json')) as f: s=json.load(f)
import pandas as pd
profiles=pd.read_csv(os.path.join(OUT,'cluster_profiles.csv'))
recs=pd.read_csv(os.path.join(OUT,'sample_recommendations.csv'))

# ---------- DOCX report ----------
doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
styles=doc.styles; styles['Normal'].font.name='Arial'; styles['Normal'].font.size=Pt(10)
for st in ['Title','Heading 1','Heading 2']:
    styles[st].font.name='Arial'

t=doc.add_paragraph(); t.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=t.add_run("SPOTIFY SONGS' GENRE SEGMENTATION\nAND RECOMMENDATION SYSTEM"); r.bold=True; r.font.size=Pt(22)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Major Project Report').bold=True
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Machine Learning / Data Analytics Project')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Dataset: Spotify playlist tracks supplied for the project')
doc.add_page_break()

sections=[
('1. Abstract', 'This project develops an automated music analysis and recommendation pipeline using Spotify playlist track data. The system performs data preprocessing, exploratory data analysis, correlation analysis, unsupervised K-Means segmentation, supervised genre prediction and nearest-neighbor recommendation. Audio characteristics such as danceability, energy, loudness, speechiness, acousticness, instrumentalness, liveness, valence, tempo, duration and popularity are used to represent songs numerically. The final system creates six interpretable music segments and provides a recommendation engine that retrieves acoustically similar tracks. A Random Forest baseline is also trained to predict the playlist genre from audio features.'),
('2. Introduction', 'Music streaming platforms contain large collections of songs described by metadata and numerical audio characteristics. A recommendation system must identify similarities between tracks and understand user-oriented groupings such as genres, moods and playlist styles. This project applies machine learning to transform raw Spotify playlist data into meaningful song segments and a simple recommendation system.'),
('3. Problem Statement', 'Manual organization of thousands of songs into meaningful groups is time-consuming and subjective. The objective is to create an automated system that analyzes song attributes, discovers groups of similar tracks, identifies genre patterns and produces recommendations based on audio-feature similarity.'),
('4. Objectives', '• Perform data preprocessing and quality checks.\n• Analyze distributions and relationships among song/audio features.\n• Generate a numerical correlation matrix.\n• Study playlist genres, subgenres and playlist names.\n• Segment songs using K-Means clustering.\n• Evaluate candidate cluster counts using silhouette, Calinski-Harabasz and Davies-Bouldin measures.\n• Build a genre prediction baseline using Random Forest.\n• Build a nearest-neighbor recommendation component.'),
('5. Dataset Description', f"The supplied CSV contains {s['input_rows']:,} rows and {s['input_columns']} columns. After preprocessing, {s['clean_rows']:,} rows remained. The dataset contains {s['genres']} playlist genres, {s['subgenres']} playlist subgenres and {s['artists']:,} unique artists. Important variables include track metadata, playlist metadata and audio features such as danceability, energy, loudness, speechiness, acousticness, instrumentalness, liveness, valence and tempo. The mentor-provided Kaggle reference is https://www.kaggle.com/datasets/maharshipandya/-spotify-tracks-dataset/code.'"),
('6. Data Preprocessing', 'The pipeline removes exact duplicate rows, fills missing text metadata with “Unknown”, converts numerical columns to numeric types, removes rows missing essential modeling fields, converts duration from milliseconds to minutes, clips bounded audio features to valid ranges, and restricts popularity to 0–100. The cleaned dataset is saved as outputs/cleaned_spotify_tracks.csv.'),
('7. Exploratory Data Analysis', 'The analysis includes popularity distribution, genre and subgenre counts, audio-feature distributions, correlation matrix, energy-versus-danceability scatter analysis, average popularity by genre, genre feature heatmap, duration distribution and popularity distribution by genre. These plots help identify the structure of the catalogue before machine learning.'),
('8. Methodology', 'The numerical audio features are standardized using StandardScaler. Candidate K values from 2 through 10 are evaluated on a representative sample. K-Means is then fitted to the complete standardized dataset. PCA is used only for two-dimensional visualization. A Random Forest classifier is trained as a supervised baseline for playlist genre prediction. Finally, a Nearest Neighbors model provides song recommendations from the learned audio-feature space.'),
('9. Cluster Selection and Segmentation', f"Silhouette evaluation produced the strongest purely unsupervised score at K={s['optimal_k_by_silhouette']}. However, the project requirement emphasizes playlist genre segmentation and the supplied data contains six playlist genres. Therefore, K=6 was selected for the final practical segmentation because it provides a richer and more interpretable recommendation space. The final K=6 silhouette score is {s['silhouette_score_final_k']:.4f}; Calinski-Harabasz is {s['calinski_harabasz']:.2f}; Davies-Bouldin is {s['davies_bouldin']:.3f}. This trade-off is explicitly documented rather than presenting K=6 as the mathematically optimal silhouette solution."),
('10. Supervised Genre Prediction', f"A Random Forest classifier was trained on the same standardized audio-feature representation to predict the six playlist genres. On a stratified 20% test split, the model achieved an accuracy of {s['genre_classifier_accuracy']*100:.2f}%. The detailed precision, recall and F1 scores are stored in outputs/classification_report.txt. This model can act as a genre-aware layer before recommendations are presented."),
('11. Recommendation System', 'The recommendation engine represents every track as a standardized feature vector. For a selected or newly specified song profile, the vector is transformed with the saved scaler and compared with the stored track vectors using Euclidean nearest-neighbor search. The system predicts a playlist genre, identifies the learned segment and returns acoustically similar songs. This is a content-based recommender and does not claim to reproduce Spotify’s proprietary production recommendation algorithm.'),
('12. Results', f"The pipeline successfully processed {s['clean_rows']:,} tracks. The final segmentation contains six clusters with distinct audio profiles. The largest segment contains 10,899 tracks and is dominated by Latin playlists; another segment is dominated by EDM and has very high instrumentalness; another is dominated by R&B with higher acousticness and lower energy; and a rap-dominant segment shows relatively high speechiness. The complete cluster profile table is included in the outputs folder."),
('13. Limitations', 'The supplied data represents playlist/track snapshots rather than a live user-history system. Playlist genre is a label attached to the source data and may not perfectly correspond to the acoustic characteristics of every song. K-Means assumes spherical cluster structure and Euclidean distance may not represent every musical similarity. The Random Forest accuracy is a baseline, not a production-grade genre classifier. Recommendations are based on audio similarity and do not use collaborative listening history.'),
('14. Future Enhancements', 'Future versions can incorporate user listening history, collaborative filtering, hybrid content-plus-collaborative recommendation, neural embeddings, automatic mood detection, artist/album embeddings, Spotify API integration, a feedback loop, personalized playlists and online model monitoring.'),
('15. Conclusion', 'The project demonstrates a complete machine-learning workflow from raw Spotify playlist data to an operational recommendation prototype. Data preprocessing and EDA establish the dataset structure, K-Means creates interpretable song segments, Random Forest provides a genre prediction baseline, and nearest-neighbor search converts the learned representation into recommendations. The result is suitable as an academic major project foundation and can be extended into a more personalized recommendation platform.')]
for h,body in sections:
    doc.add_heading(h,level=1)
    for para in body.split('\n'):
        if para.startswith('• '): doc.add_paragraph(para[2:],style='List Bullet')
        else: doc.add_paragraph(para)

# figures section

doc.add_heading('16. Key Visualizations',level=1)
figs=['01_popularity_distribution.png','02_top_genres_count.png','03_top_subgenres_count.png','04_audio_feature_distributions.png','05_correlation_matrix.png','06_energy_vs_danceability.png','07_avg_popularity_genre.png','08_genre_feature_heatmap.png','09_duration_distribution.png','10_popularity_by_genre_boxplot.png','11_silhouette_scores.png','12_pca_clusters.png','13_cluster_sizes.png','14_genre_classification_confusion_matrix.png','15_feature_importance.png']
for i,f in enumerate(figs,1):
    doc.add_paragraph(f'Figure {i}: {f.replace("_"," ").replace(".png","").title()}')
    doc.add_picture(os.path.join(PLOT,f),width=Inches(6.3)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER

# cluster table

doc.add_heading('17. Final Cluster Profiles',level=1)
table=doc.add_table(rows=1,cols=6); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style='Table Grid'
for c,h in enumerate(['Cluster','Tracks','Dominant Genre','Energy','Danceability','Popularity']): table.rows[0].cells[c].text=h
for _,r in profiles.iterrows():
    cells=table.add_row().cells
    vals=[int(r.cluster),int(r.track_count),str(r.playlist_genre_dominant),f"{r.energy:.3f}",f"{r.danceability:.3f}",f"{r.track_popularity:.1f}"]
    for c,v in enumerate(vals): cells[c].text=str(v)

doc.add_heading('18. Sample Recommendations',level=1)
t=doc.add_table(rows=1,cols=5); t.style='Table Grid'
for c,h in enumerate(['Track','Artist','Genre','Subgenre','Segment']): t.rows[0].cells[c].text=h
for _,r in recs.head(8).iterrows():
    cells=t.add_row().cells
    for c,v in enumerate([r.track_name,r.track_artist,r.playlist_genre,r.playlist_subgenre,r.segment_name]): cells[c].text=str(v)

path=os.path.join(DOCS,'Major_Project_Report.docx'); doc.save(path)

# ---------- PPTX ----------
prs=Presentation(); prs.slide_width=PInches(13.333); prs.slide_height=PInches(7.5)
def add_slide(title, bullets=None, image=None):
    layout=prs.slide_layouts[5] if image else prs.slide_layouts[1]
    slide=prs.slides.add_slide(layout); slide.shapes.title.text=title
    if bullets:
        box=slide.placeholders[1].text_frame; box.clear()
        for i,b in enumerate(bullets):
            p=box.paragraphs[0] if i==0 else box.add_paragraph(); p.text=b; p.font.size=PPt(22); p.space_after=PPt(8)
    if image:
        slide.shapes.add_picture(image,PInches(1.0),PInches(1.55),width=PInches(11.3),height=PInches(5.4))
    return slide

slide=prs.slides.add_slide(prs.slide_layouts[0]); slide.shapes.title.text="Spotify Songs' Genre Segmentation"; slide.placeholders[1].text='Major Project\nMachine Learning & Recommendation System'
add_slide('Problem Statement',['Large music catalogues are difficult to organize manually.','We need automated song segmentation using audio characteristics.','The final system should support genre understanding and recommendations.'])
add_slide('Objectives',['Preprocess and clean the Spotify dataset.','Perform comprehensive EDA and correlation analysis.','Discover song segments using K-Means clustering.','Predict playlist genre using a machine-learning baseline.','Recommend similar tracks using nearest-neighbor search.'])
add_slide('Dataset',['32,833 supplied track records','23 columns','6 playlist genres and 24 subgenres','10,693 unique artists','Audio features: danceability, energy, loudness, speechiness, acousticness, instrumentalness, liveness, valence, tempo and duration'])
add_slide('Data Preprocessing',['Remove exact duplicates','Handle missing track metadata','Convert numeric fields to valid numeric types','Convert duration from milliseconds to minutes','Clip bounded audio features and popularity to valid ranges'])
add_slide('EDA – Popularity',image=os.path.join(PLOT,'01_popularity_distribution.png'))
add_slide('EDA – Correlation Matrix',image=os.path.join(PLOT,'05_correlation_matrix.png'))
add_slide('EDA – Genre Feature Profiles',image=os.path.join(PLOT,'08_genre_feature_heatmap.png'))
add_slide('Clustering Methodology',['Standardize numerical audio features.','Evaluate K=2 to K=10 using silhouette, Calinski-Harabasz and Davies-Bouldin scores.','Use K=6 for the final practical genre-segmentation space.','Use PCA only for visualization of the learned clusters.'])
add_slide('Cluster Evaluation',image=os.path.join(PLOT,'11_silhouette_scores.png'))
add_slide('Final Music Segments',image=os.path.join(PLOT,'12_pca_clusters.png'))
add_slide('Genre Prediction',['Random Forest classifier trained on audio features.','Stratified 80/20 train-test split.','Accuracy: {:.2f}%'.format(s['genre_classifier_accuracy']*100),'Best-performing classes include EDM and Rock; Pop is more difficult to distinguish from neighboring styles.'])
add_slide('Recommendation System',['Input: a song/audio-feature profile.','Scale the feature vector using the saved scaler.','Predict playlist genre and K-Means segment.','Retrieve nearest tracks using Euclidean distance.','Return track, artist, genre, subgenre and segment information.'])
add_slide('Results',['6 final music segments','K=6 silhouette score: {:.4f}'.format(s['silhouette_score_final_k']),'Genre classification accuracy: {:.2f}%'.format(s['genre_classifier_accuracy']*100),'Working recommendation prototype with saved models and Streamlit UI'])
add_slide('Conclusion & Future Work',['Complete end-to-end ML pipeline achieved.','Content-based recommendations are functional.','Future: collaborative filtering, user profiles, embeddings, hybrid recommendation and live API integration.'])
add_slide('Thank You',['Questions?'])
pp=os.path.join(DOCS,'Major_Project_Presentation.pptx'); prs.save(pp)
print(path); print(pp)
